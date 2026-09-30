"""Pure REST API backend for the Job Hunter pipeline."""

import os
import sys
import time
import shutil
import threading
import subprocess
import sqlite3
from pathlib import Path

from flask import Flask, jsonify, request, send_file, send_from_directory
from openpyxl import load_workbook


HERE = Path(__file__).parent.resolve()
PKG = HERE / "jobseeker"
if not PKG.exists() and (HERE / "main.py").exists():
	PKG = HERE

EXCEL_PATH = PKG / "data" / "applied.xlsx"
DB_PATH = PKG / "data" / "jobs.db"
ROOT = PKG.parent if PKG.name == "jobseeker" else PKG
STATIC_DIR = HERE / "static"
INDEX_HTML = HERE / "index.html"
if not INDEX_HTML.exists() and (HERE / "frontend" / "index.html").exists():
	INDEX_HTML = HERE / "frontend" / "index.html"
if not STATIC_DIR.exists() and (HERE / "frontend" / "static").exists():
	STATIC_DIR = HERE / "frontend" / "static"
VERSION = "1.0.0"

STATE = {
	"running": False,
	"started_at": None,
	"finished_at": None,
	"exit_code": None,
	"logs": [],
	"excel_ready": False,
	"last_error": None,
	"pid": None,
}
LOCK = threading.Lock()
MAX_LOGS = 8000


def _run_pipeline():
	with LOCK:
		STATE.update({
			"running": True,
			"started_at": time.time(),
			"finished_at": None,
			"exit_code": None,
			"logs": [],
			"excel_ready": False,
			"last_error": None,
			"pid": None,
		})

	env = os.environ.copy()
	env["PYTHONUNBUFFERED"] = "1"
	env["PYTHONIOENCODING"] = "utf-8"
	env["PYTHONPATH"] = str(ROOT) + os.pathsep + env.get("PYTHONPATH", "")
	if PKG.name == "jobseeker":
		cmd = [sys.executable, "-m", "jobseeker.main"]
		cwd = str(ROOT)
	else:
		cmd = [sys.executable, "main.py"]
		cwd = str(PKG)

	try:
		proc = subprocess.Popen(
			cmd,
			cwd=cwd,
			stdout=subprocess.PIPE,
			stderr=subprocess.STDOUT,
			text=True,
			bufsize=1,
			encoding="utf-8",
			errors="replace",
			env=env,
		)
	except Exception as exc:
		with LOCK:
			STATE["running"] = False
			STATE["finished_at"] = time.time()
			STATE["last_error"] = f"failed to start subprocess: {exc}"
		return

	with LOCK:
		STATE["pid"] = proc.pid

	try:
		for raw in iter(proc.stdout.readline, ""):
			with LOCK:
				STATE["logs"].append(raw.rstrip("\n"))
				if len(STATE["logs"]) > MAX_LOGS:
					STATE["logs"] = STATE["logs"][-MAX_LOGS:]
	except Exception as exc:
		with LOCK:
			STATE["logs"].append(f"[server] read error: {exc}")

	proc.wait()
	with LOCK:
		STATE["running"] = False
		STATE["finished_at"] = time.time()
		STATE["exit_code"] = proc.returncode
		STATE["pid"] = None
		STATE["excel_ready"] = EXCEL_PATH.exists()
		if proc.returncode != 0:
			STATE["last_error"] = f"pipeline exited with code {proc.returncode}"


app = Flask(__name__)


@app.after_request
def _cors(response):
	response.headers["Access-Control-Allow-Origin"] = "*"
	response.headers["Access-Control-Allow-Methods"] = "GET, POST, OPTIONS"
	response.headers["Access-Control-Allow-Headers"] = "Content-Type, Authorization"
	response.headers["Access-Control-Max-Age"] = "3600"
	return response


@app.route("/api/<path:_any>", methods=["OPTIONS"])
def _preflight(_any):
	return ("", 204)


@app.route("/")
def root():
	if INDEX_HTML.exists():
		return send_file(str(INDEX_HTML))
	return f"Job Hunter API v{VERSION}\n", 200, {"Content-Type": "text/plain; charset=utf-8"}


if STATIC_DIR.exists():
	@app.route("/static/<path:filename>")
	def static_files(filename):
		return send_from_directory(str(STATIC_DIR), filename)


@app.route("/api/health")
def api_health():
	return jsonify({
		"ok": True,
		"version": VERSION,
		"root": str(ROOT),
		"package": str(PKG),
		"excel_path": str(EXCEL_PATH),
		"db_path": str(DB_PATH),
		"index_html_exists": INDEX_HTML.exists(),
	})


@app.route("/api/status")
def api_status():
	excel_exists = EXCEL_PATH.exists()
	with LOCK:
		payload = dict(STATE)
		payload["excel_ready"] = STATE["excel_ready"] or excel_exists
		payload["log_count"] = len(STATE["logs"])
	payload["excel_size"] = EXCEL_PATH.stat().st_size if excel_exists else 0
	payload["excel_mtime"] = EXCEL_PATH.stat().st_mtime if excel_exists else None
	payload.pop("logs", None)
	return jsonify(payload)


@app.route("/api/logs")
def api_logs():
	try:
		after = max(0, int(request.args.get("after", 0)))
	except ValueError:
		after = 0
	try:
		limit = max(0, int(request.args.get("limit", 0)))
	except ValueError:
		limit = 0
	with LOCK:
		total = len(STATE["logs"])
		chunk = STATE["logs"][after:]
		if limit:
			chunk = chunk[:limit]
	return jsonify({"logs": chunk, "total": total, "after": after})


def _open_workbook():
	return load_workbook(str(EXCEL_PATH), read_only=True, data_only=True)


@app.route("/api/sheets")
def api_sheets():
	if not EXCEL_PATH.exists():
		return jsonify({"error": "Excel not ready", "sheets": []}), 404
	try:
		wb = _open_workbook()
		sheets = list(wb.sheetnames)
		wb.close()
		return jsonify({"sheets": sheets})
	except Exception as exc:
		return jsonify({"error": str(exc), "sheets": []}), 500


@app.route("/api/results")
def api_results():
	if not EXCEL_PATH.exists():
		return jsonify({"error": "Excel not ready", "columns": [], "rows": []}), 404
	try:
		wb = _open_workbook()
		sheet_name = request.args.get("sheet")
		if sheet_name not in wb.sheetnames:
			sheet_name = wb.sheetnames[0]
		ws = wb[sheet_name]
		rows_iter = ws.iter_rows(values_only=True)
		header = next(rows_iter, ())
		columns = [str(value) if value is not None else "" for value in header]
		rows = [
			{col: (row[i] if i < len(row) else None) for i, col in enumerate(columns)}
			for row in rows_iter
		]
		wb.close()
		return jsonify({"sheet": sheet_name, "columns": columns, "rows": rows})
	except Exception as exc:
		return jsonify({"error": str(exc), "columns": [], "rows": []}), 500


@app.route("/api/jobs")
def api_jobs():
	if not DB_PATH.exists():
		return jsonify({"error": "DB not found", "rows": [], "count": 0}), 404
	try:
		limit = max(1, int(request.args.get("limit", 200)))
	except ValueError:
		limit = 200
	try:
		min_score = int(request.args.get("min_score", 0))
	except ValueError:
		min_score = 0
	status = request.args.get("status")
	sql = "SELECT * FROM jobs WHERE COALESCE(match_score, 0) >= ?"
	args = [min_score]
	if status:
		sql += " AND status = ?"
		args.append(status)
	sql += " ORDER BY match_score DESC LIMIT ?"
	args.append(limit)
	try:
		conn = sqlite3.connect(str(DB_PATH))
		conn.row_factory = sqlite3.Row
		rows = [dict(row) for row in conn.execute(sql, args).fetchall()]
		conn.close()
		return jsonify({"rows": rows, "count": len(rows)})
	except Exception as exc:
		return jsonify({"error": str(exc), "rows": [], "count": 0}), 500


@app.route("/api/run", methods=["POST"])
def api_run():
	with LOCK:
		if STATE["running"]:
			return jsonify({"error": "Pipeline already running"}), 409
	threading.Thread(target=_run_pipeline, daemon=True).start()
	return jsonify({"ok": True, "message": "Pipeline started"})


@app.route("/api/stop", methods=["POST"])
def api_stop():
	with LOCK:
		pid = STATE["pid"]
	if not pid:
		return jsonify({"error": "Not running"}), 400
	try:
		if sys.platform.startswith("win"):
			subprocess.run(["taskkill", "/F", "/T", "/PID", str(pid)], capture_output=True)
		else:
			import signal
			os.kill(pid, signal.SIGTERM)
		return jsonify({"ok": True, "message": "Stop signal sent"})
	except Exception as exc:
		return jsonify({"error": str(exc)}), 500


@app.route("/api/download")
def api_download():
	if not EXCEL_PATH.exists():
		return jsonify({"error": "Excel not ready yet"}), 404
	return send_file(
		str(EXCEL_PATH),
		as_attachment=True,
		download_name="applied.xlsx",
		mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
	)


@app.route("/api/clear", methods=["POST"])
def api_clear():
	"""Delete generated data and reset the in-memory server state."""
	with LOCK:
		if STATE["running"]:
			return jsonify({"error": "Stop the pipeline first"}), 409

	deleted = []
	targets = [
		DB_PATH,
		EXCEL_PATH,
		PKG / "data" / "cold_outreach.xlsx",
		PKG / "data" / "profile.json",
	]
	for path in targets:
		try:
			if path.exists():
				path.unlink()
				deleted.append(path.name)
		except Exception as exc:
			return jsonify({"error": f"failed to delete {path.name}: {exc}"}), 500

	with LOCK:
		STATE["logs"] = []
		STATE["excel_ready"] = False
		STATE["exit_code"] = None
		STATE["started_at"] = None
		STATE["finished_at"] = None
		STATE["last_error"] = None

	return jsonify({"ok": True, "deleted": deleted})


if __name__ == "__main__":
	(PKG / "data").mkdir(parents=True, exist_ok=True)
	(PKG / "__init__.py").touch(exist_ok=True)
	root_cfg = HERE / "config.json"
	pkg_cfg = PKG / "config.json"
	if root_cfg.exists() and not pkg_cfg.exists():
		try:
			shutil.copy(root_cfg, pkg_cfg)
			print(f"[server] copied config.json → {pkg_cfg}")
		except Exception as exc:
			print(f"[server] warn: could not copy config.json: {exc}")

	print("=" * 60)
	print(f"JOB HUNTER API v{VERSION}")
	print("=" * 60)
	print(f"  root:      {ROOT}")
	print(f"  package:   {PKG}")
	print(f"  excel:     {EXCEL_PATH}")
	print(f"  db:        {DB_PATH}")
	print(f"  python:    {sys.executable}")
	print(f"  index:     {INDEX_HTML if INDEX_HTML.exists() else '(not yet)'}")
	print("  listening: http://127.0.0.1:5000")
	print("=" * 60)
	app.run(host="127.0.0.1", port=5000, debug=False, threaded=True)

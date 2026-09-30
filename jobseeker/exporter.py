# exporter.py
"""Excel-first exporter: sorted, hyperlinked, color-coded, multi-sheet."""
import json
import re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ─── Columns ────────────────────────────────────────────────
COLS = [
    "rank", "match_score", "company", "title", "location",
    "apply_link", "job_url",
    "salary_min", "salary_max", "currency",
    "matched_skills", "reason", "source",
    "contact_email", "contact_phone", "company_email", "contact_page",
    "status", "notes",
]

HEADERS = {
    "rank": "#",
    "match_score": "Score",
    "company": "Company",
    "title": "Role",
    "location": "Location",
    "apply_link": "Apply Now",
    "job_url": "Job URL",
    "salary_min": "Sal Min",
    "salary_max": "Sal Max",
    "currency": "Cur",
    "matched_skills": "Matched Skills",
    "reason": "Why",
    "source": "Source",
    "contact_email": "Email",
    "contact_phone": "Phone",
    "company_email": "Company Email",
    "contact_page": "Contact Page",
    "status": "Status",
    "notes": "Notes",
}

COL_WIDTHS = {
    "rank": 5, "match_score": 7, "company": 22, "title": 42,
    "location": 22, "apply_link": 14, "job_url": 45,
    "salary_min": 9, "salary_max": 9, "currency": 6,
    "matched_skills": 40, "reason": 50, "source": 16,
    "contact_email": 28, "contact_phone": 18, "company_email": 30,
    "contact_page": 30, "status": 12, "notes": 20,
}


def _pretty(v):
    if not v:
        return ""
    if isinstance(v, list):
        return ", ".join(str(x) for x in v)
    try:
        parsed = json.loads(v)
        if isinstance(parsed, list):
            return ", ".join(str(x) for x in parsed)
        return str(parsed)
    except Exception:
        return str(v)


def _score_color(score):
    if score is None:
        return None
    if score >= 90:
        return "1F7A3A"
    if score >= 80:
        return "9CCC65"
    if score >= 70:
        return "FFD54F"
    if score >= 60:
        return "FFAB40"
    return "EF9A9A"


def _write_sheet(wb, name, rows):
    ws = wb.create_sheet(name)

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1A237E")
    for ci, col in enumerate(COLS, start=1):
        c = ws.cell(row=1, column=ci, value=HEADERS.get(col, col))
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    for ri, row in enumerate(rows, start=2):
        row["rank"] = ri - 1
        score = row.get("match_score") or 0
        color = _score_color(score)
        for ci, col in enumerate(COLS, start=1):
            val = row.get(col)
            if col == "matched_skills":
                val = _pretty(val)

            cell = ws.cell(row=ri, column=ci, value=val)

            if col == "match_score" and color:
                cell.fill = PatternFill("solid", fgColor=color)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")

            if col == "apply_link":
                url = row.get("apply_url") or row.get("job_url")
                if url:
                    cell.value = "▶ Apply"
                    cell.hyperlink = url
                    cell.style = "Hyperlink"
                    cell.font = Font(color="0563C1", underline="single", bold=True)
                    cell.alignment = Alignment(horizontal="center")

            if col == "job_url" and val:
                cell.hyperlink = val
                cell.font = Font(color="0563C1", underline="single")

            if col == "contact_email" and val:
                cell.hyperlink = f"mailto:{val}"
                cell.font = Font(color="0563C1", underline="single")

            if col == "contact_phone" and val:
                cell.hyperlink = f"tel:{re.sub(r'[^+0-9]', '', str(val))}"
                cell.font = Font(color="0563C1", underline="single")

            if col == "company_email" and val:
                cell.hyperlink = f"mailto:{val}"
                cell.font = Font(color="0563C1", underline="single")

            if col == "contact_page" and val:
                cell.hyperlink = val
                cell.font = Font(color="0563C1", underline="single")

            if col in ("matched_skills", "reason", "title"):
                cell.alignment = Alignment(vertical="top", wrap_text=False)
            else:
                cell.alignment = Alignment(vertical="top")

    for ci, col in enumerate(COLS, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = COL_WIDTHS.get(col, 18)

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{max(1, len(rows) + 1)}"
    return ws


def _split_cold_email(raw):
    if not raw:
        return "", ""
    try:
        d = json.loads(raw) if isinstance(raw, str) else raw
        return d.get("subject", ""), d.get("body", "")
    except Exception:
        return "", ""


def _guess_email(company):
    if not company:
        return ""
    return f"careers@{re.sub(r'[^a-z0-9]', '', company.lower())}.com"


def _row_from_job(j):
    return {
        "match_score": j.get("match_score") or 0,
        "company": j.get("company"),
        "title": j.get("title"),
        "location": j.get("location"),
        "apply_url": j.get("apply_url") or j.get("url"),
        "job_url": j.get("url"),
        "salary_min": j.get("salary_min"),
        "salary_max": j.get("salary_max"),
        "currency": j.get("currency"),
        "matched_skills": j.get("matched_skills"),
        "reason": j.get("reason"),
        "source": j.get("source"),
        "contact_email": j.get("contact_email") or _guess_email(j.get("company")),
        "contact_phone": j.get("contact_phone") or j.get("company_phone"),
        "company_email": j.get("company_email"),
        "contact_page": j.get("contact_page") or j.get("apply_url"),
        "status": j.get("status"),
        "notes": j.get("notes"),
    }


def _is_applied(job):
    return (job.get("status") or "").lower() == "applied"


def export_all(store, out_path, min_score=60, limit=1500):
    """Build a workbook with active jobs and a separate applied history sheet."""
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)

    scored = [j for j in store.all() if (j.get("match_score") or 0) >= min_score]
    scored = scored[:limit]
    scored.sort(key=lambda j: -(j.get("match_score") or 0))
    not_applied = [j for j in scored if not _is_applied(j)]
    applied = [j for j in scored if _is_applied(j)]

    top = [j for j in not_applied if (j.get("match_score") or 0) >= 75]
    auto = [j for j in not_applied if (j.get("match_score") or 0) >= 80]
    south_kw = [
        "kochi", "cochin", "ernakulam", "trivandrum", "thiruvananthapuram",
        "kozhikode", "calicut", "thrissur", "kollam", "kottayam",
        "bangalore", "bengaluru", "chennai", "hyderabad", "coimbatore",
        "mysore", "mysuru", "mangalore", "mangaluru", "madurai",
        "trichy", "tiruchirappalli", "vijayawada", "visakhapatnam",
        "vizag", "puducherry", "pondicherry", "salem", "tirupati", "nellore",
    ]
    south = [j for j in not_applied if any(k in (j.get("location") or "").lower() for k in south_kw)]

    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "Apply First (80+)", [_row_from_job(j) for j in auto])
    _write_sheet(wb, "Top Matches (75+)", [_row_from_job(j) for j in top])
    _write_sheet(wb, "South India", [_row_from_job(j) for j in south])
    _write_sheet(wb, "All Jobs", [_row_from_job(j) for j in not_applied])
    _write_sheet(wb, "Applied Jobs", [_row_from_job(j) for j in applied])
    wb.save(out_path)

    print(f"[export] {out_path}")
    print(f"         Apply First (80+): {len(auto)}")
    print(f"         Top Matches (75+): {len(top)}")
    print(f"         South India:       {len(south)}")
    print(f"         All Jobs:          {len(not_applied)}")
    print(f"         Applied Jobs:      {len(applied)}")
    return out_path
